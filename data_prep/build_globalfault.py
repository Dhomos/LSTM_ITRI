# -*- coding: utf-8 -*-
"""
冰機故障標記程式 (GlobalFault labeler)
=====================================
將一批 alarm*.xlsx 檔案（欄位為 Date + 多個 0/1 警報欄位）整理成
Date / GlobalFault 兩欄的標記結果，供後續冰機預警模型使用。

規則：
  - 每一列資料中，除了被排除的欄位（預設為「備用-M61」）以外，
    只要任一欄位出現非 0 的值，該列 GlobalFault = 1，否則為 0。
  - 為了讓輸出檔案容易開啟，當總列數超過 --max-rows-per-file 時，
    會自動依「日期」切成多個檔案 (每個檔案內部不會把同一天拆散)，
    檔名會附上該檔案涵蓋的日期範圍。若資料量夠小則只會產生一個檔案。

用法：
    py -3 build_globalfault.py --input-dir "E:\\Downloads\\你的場域冰機資料"

    # 自訂檔名 pattern（預設 alarm*.xlsx，適用任何地點命名，如 alarmXX機房2024...）
    py -3 build_globalfault.py --input-dir "路徑" --pattern "alarm*.xlsx"

    # 自訂輸出資料夾與檔名前綴（預設輸出到 input-dir，前綴為資料夾名稱）
    py -3 build_globalfault.py --input-dir "路徑" --output-dir "路徑2" --output-prefix "我的冰機"

    # 自訂每個檔案最多幾列 (預設 200000，約可對應 4~5 個月的每分鐘資料)
    py -3 build_globalfault.py --input-dir "路徑" --max-rows-per-file 100000

    # 若排除欄位名稱不同，可用逗號分隔指定多個
    py -3 build_globalfault.py --input-dir "路徑" --exclude-cols "備用-M61,備用2"

這支程式只負責「標記故障」。接著要用 merge_site_dataset.py，
把這裡輸出的 GlobalFault 標記，跟 analog 感測器數值合併成模型要用的單一CSV。
"""
import argparse
import datetime
import glob
import os
import sys

import openpyxl
from openpyxl.cell import WriteOnlyCell

EXCEL_HARD_ROW_LIMIT = 1_048_575  # Excel 單一工作表列數上限 (扣除表頭)
DEFAULT_MAX_ROWS_PER_FILE = 200_000


def parse_date(v):
    if isinstance(v, datetime.datetime):
        return v
    s = str(v).strip()
    try:
        return datetime.datetime.fromisoformat(s)
    except ValueError:
        return datetime.datetime.strptime(s, '%Y-%m-%d %H:%M:%S')


def is_nonzero(v):
    if v is None:
        return False
    if isinstance(v, str):
        s = v.strip()
        if s == '':
            return False
        try:
            return float(s) != 0
        except ValueError:
            return True  # 非數字內容，視為異常標記
    try:
        return float(v) != 0
    except (TypeError, ValueError):
        return bool(v)


class ChunkedFileWriter:
    """
    依日期切檔：當目前檔案列數達到 max_rows_per_file，且下一列的日期
    和目前檔案最後一列的日期不同 (不會把同一天拆到兩個檔案) 時，
    就關閉目前檔案、另開一個新檔案。
    """

    def __init__(self, output_dir, prefix, max_rows_per_file):
        self.output_dir = output_dir
        self.prefix = prefix
        self.max_rows_per_file = min(max_rows_per_file, EXCEL_HARD_ROW_LIMIT)

        self.wb = None
        self.ws = None
        self.rows_in_file = 0
        self.chunk_start_date = None
        self.last_date = None
        self.saved_files = []

    def _open_new_file(self, start_date):
        self.wb = openpyxl.Workbook(write_only=True)
        self.ws = self.wb.create_sheet(title='GlobalFault')
        self.ws.append(('Date', 'GlobalFault'))
        self.rows_in_file = 0
        self.chunk_start_date = start_date

    def _close_current_file(self):
        if self.wb is None or self.rows_in_file == 0:
            return
        start_str = self.chunk_start_date.strftime('%Y%m%d')
        end_str = self.last_date.strftime('%Y%m%d')
        if start_str == end_str:
            filename = f'{self.prefix}_{start_str}.xlsx'
        else:
            filename = f'{self.prefix}_{start_str}~{end_str}.xlsx'
        out_path = os.path.join(self.output_dir, filename)
        self.wb.save(out_path)
        self.saved_files.append((out_path, self.rows_in_file))
        print(f'  已儲存 {filename}: {self.rows_in_file} 列')

    def append(self, date_val, fault):
        if self.wb is None:
            self._open_new_file(date_val)
        elif (self.rows_in_file >= self.max_rows_per_file
              and date_val.date() != self.last_date.date()):
            self._close_current_file()
            self._open_new_file(date_val)

        date_cell = WriteOnlyCell(self.ws, value=date_val)
        date_cell.number_format = 'yyyy-mm-dd hh:mm:ss'
        self.ws.append((date_cell, fault))
        self.rows_in_file += 1
        self.last_date = date_val

    def close(self):
        self._close_current_file()


def iter_globalfault_rows(path, exclude_cols, warn_missing_exclude=True):
    """對單一 alarm*.xlsx 檔案，逐列yield (date_val, fault)。

    抽成獨立函式是為了讓 merge_site_dataset.py 可以直接重用同一套判斷邏輯，
    不用在兩支程式裡各寫一次、容易兩邊邏輯兜不起來。
    """
    wb = openpyxl.load_workbook(path, read_only=True, data_only=True)
    try:
        ws_in = wb[wb.sheetnames[0]]
        row_iter = ws_in.iter_rows(values_only=True)
        header = next(row_iter)

        exclude_idx = set()
        for col_name in exclude_cols:
            if col_name in header:
                exclude_idx.add(header.index(col_name))
            elif warn_missing_exclude:
                print(f'  [警告] 找不到排除欄位「{col_name}」於檔案 {os.path.basename(path)}，將忽略此設定')

        check_indices = [i for i in range(1, len(header)) if i not in exclude_idx]

        for row in row_iter:
            if row[0] is None:
                continue
            date_val = parse_date(row[0])
            fault = 1 if any(is_nonzero(row[i]) for i in check_indices) else 0
            yield date_val, fault
    finally:
        wb.close()


def process(input_dir, pattern, output_dir, prefix, max_rows_per_file, exclude_cols):
    files = sorted(glob.glob(os.path.join(input_dir, pattern)))
    if not files:
        print(f'找不到符合條件的檔案：{os.path.join(input_dir, pattern)}')
        sys.exit(1)

    os.makedirs(output_dir, exist_ok=True)

    print('將依序處理下列檔案：')
    for f in files:
        print(' ', f)

    writer = ChunkedFileWriter(output_dir, prefix, max_rows_per_file)

    total_rows = 0
    total_fault = 0

    for path in files:
        file_rows = 0
        for date_val, fault in iter_globalfault_rows(path, exclude_cols):
            writer.append(date_val, fault)
            file_rows += 1
            total_rows += 1
            total_fault += fault

        print(f'已處理 {os.path.basename(path)}：{file_rows} 列')

    writer.close()

    print('-' * 40)
    print('總列數:', total_rows)
    print('故障列數 (GlobalFault=1):', total_fault)
    print(f'輸出檔案數: {len(writer.saved_files)}')
    for out_path, cnt in writer.saved_files:
        print(f'  {out_path} ({cnt} 列)')


def main():
    parser = argparse.ArgumentParser(description='冰機 alarm 資料 GlobalFault 標記程式')
    parser.add_argument('--input-dir', default=None, help='存放 alarm*.xlsx 檔案的資料夾')
    parser.add_argument('--pattern', default='alarm*.xlsx', help='檔名比對樣式 (預設: alarm*.xlsx)')
    parser.add_argument('--output-dir', default=None, help='輸出資料夾 (預設與 input-dir 相同)')
    parser.add_argument('--output-prefix', default=None, help='輸出檔名前綴 (預設為資料夾名稱_GlobalFault)')
    parser.add_argument('--max-rows-per-file', type=int, default=DEFAULT_MAX_ROWS_PER_FILE,
                         help=f'每個輸出檔案最多幾列，超過會依日期自動切成多檔 (預設: {DEFAULT_MAX_ROWS_PER_FILE})')
    parser.add_argument('--exclude-cols', default='備用-M61',
                         help='要排除於故障判斷之外的欄位名稱，逗號分隔 (預設: 備用-M61)')
    args = parser.parse_args()

    input_dir = args.input_dir
    if not input_dir:
        # 沒有給參數時，直接互動詢問資料夾路徑，方便直接雙擊執行
        input_dir = input('請輸入存放 alarm*.xlsx 檔案的資料夾路徑: ').strip().strip('"')

    if not os.path.isdir(input_dir):
        print(f'找不到資料夾：{input_dir}')
        sys.exit(1)

    exclude_cols = [c.strip() for c in args.exclude_cols.split(',') if c.strip()]

    output_dir = args.output_dir or input_dir
    if args.output_prefix:
        prefix = args.output_prefix
    else:
        folder_name = os.path.basename(os.path.normpath(input_dir))
        prefix = f'{folder_name}_GlobalFault'

    process(input_dir, args.pattern, output_dir, prefix, args.max_rows_per_file, exclude_cols)


if __name__ == '__main__':
    try:
        main()
    except Exception as e:
        print('發生錯誤:', e)
        raise
    finally:
        if len(sys.argv) == 1:
            input('\n按 Enter 結束...')
